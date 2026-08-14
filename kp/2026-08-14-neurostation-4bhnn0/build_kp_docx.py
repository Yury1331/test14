#!/usr/bin/env python3
"""Сборка КП № КП-2026-08-14-001 в формате Microsoft Word (.docx)."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x4E, 0x79)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC0, 0x00, 0x00)

OUT = (
    "/workspace/kp/2026-08-14-neurostation-4bhnn0/"
    "КП-2026-08-14-001_NeuroStation_Astra_9800R_4BHNN0.docx"
)


def set_run_font(run, *, size=11, bold=False, color=BLACK, name="Times New Roman"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcpr.append(shd)


def set_cell_border(cell) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "BFBFBF")
        tc_borders.append(el)
    tcpr.append(tc_borders)


def set_cell_text(cell, text, *, bold=False, size=9, align="left", color=BLACK, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    if fill:
        shade_cell(cell, fill)
    set_cell_border(cell)
    cell.vertical_alignment = 1  # center


def add_heading_bar(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=13, bold=True, color=NAVY, name="Calibri")


def add_body(doc, text: str, *, size=11, italic=False, color=BLACK, space_after=6) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, color=color)
    run.italic = italic


def add_table(doc, headers, rows, col_widths=None, money_cols=None):
    money_cols = money_cols or set()
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8, align="center", color=WHITE, fill="1F4E79")
    for r_idx, row in enumerate(rows):
        fill = "F2F2F2" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            align = "right" if c_idx in money_cols else ("center" if c_idx in (0, 3, 4) else "left")
            set_cell_text(
                table.rows[r_idx + 1].cells[c_idx],
                val,
                size=8,
                align=align,
                fill=fill,
                bold=c_idx in money_cols and r_idx == len(rows) - 1 and False,
            )
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    return table


def set_narrow_margins(doc) -> None:
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)


def add_footer(doc) -> None:
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("КП-2026-08-14-001 ред. 4  ·  две заявки ЭТП ГПБ  ·  цены с НДС 22%  ·  стр. ")
    set_run_font(run, size=8, color=MUTED, name="Calibri")

    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run2 = p.add_run()
    run2._r.append(fld_char_begin)
    run2._r.append(instr)
    run2._r.append(fld_char_end)
    set_run_font(run2, size=8, color=MUTED, name="Calibri")


def build() -> None:
    doc = Document()
    set_narrow_margins(doc)
    add_footer(doc)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ")
    set_run_font(run, size=18, bold=True, color=NAVY, name="Calibri")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    run = sub.add_run("№ КП-2026-08-14-001, редакция 4 от 14 августа 2026 г.")
    set_run_font(run, size=13, bold=True, color=NAVY, name="Calibri")

    add_body(
        doc,
        "Две отдельные процедуры ЭТП ГПБ, секция «Торговый портал», режим 44-ФЗ, способ — ценовой запрос. "
        "Опубликованы 12.08.2026. Срок подачи — до 18.08.2026, 00:00 МСК. НМЦ не указана. "
        "Валюта — рубль, НДС 22% включён в цены. Заявки на площадку подаются отдельно, лоты не объединять.",
        size=11,
    )
    add_body(
        doc,
        "Ориентировочный расчёт по открытым ценам на 14.08.2026. Цена в заявке на ЭТП — ответственность участника. "
        "Карточки процедур с этой среды не открылись; состав и сроки взяты из извещения заказчика.",
        size=10,
        italic=True,
        color=MUTED,
    )

    meta = [
        ("Заказчик", "ФГКУ «Донской спасательный центр МЧС России»"),
        ("ИНН / КПП / ОГРН", "6102006605 / 610201001 / 1026100666525"),
        ("Адрес поставки", "346709, Ростовская обл., Аксайский р-н, п. Ковалевка, ул. Салютная, д. 2"),
        ("Контакт", "Пахомов И.В., donsc@dsc.61.mchs.gov.ru, +7 (86350) 2-71-98"),
        ("Поставщик", "к заполнению (DSSL / авторизованный дилер TRASSIR)"),
        ("Срок подачи", "до 18.08.2026, 00:00 МСК"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        set_cell_text(table.rows[i].cells[0], k, bold=True, size=10, fill="D6E4F0")
        set_cell_text(table.rows[i].cells[1], v, size=10)
        table.rows[i].cells[0].width = Cm(4.5)
        table.rows[i].cells[1].width = Cm(13.5)

    add_heading_bar(doc, "А. Заявка 254364 (№ ЭТП 998818) — регистраторы")
    add_table(
        doc,
        ["№", "Наименование", "Артикул", "Ед.", "Кол-во", "Цена с НДС, ₽", "Сумма с НДС, ₽"],
        [
            [
                "1",
                "IP-видеорегистратор TRASSIR NeuroStation Astra 9800R/128-S, 128 кан., Astra Linux SE «Смоленск», 8×HDD (диски не входят), 2× Neuro Detector, СТ-1",
                "NeuroStation 9800R/128-S",
                "шт.",
                "2",
                "477 740",
                "955 480",
            ],
        ],
        money_cols={5, 6},
    )
    add_table(
        doc,
        ["Показатель", "Сумма, ₽"],
        [
            ["Всего с НДС по заявке 254364", "955 480"],
            ["в том числе НДС 22%", "172 299,67"],
            ["без НДС", "783 180,33"],
        ],
        money_cols={1},
    )

    add_heading_bar(doc, "Б. Заявка 254397 (№ ЭТП 998881) — камеры 4BHNN0-0-0-0")
    add_body(
        doc,
        "Предмет извещения — 51 шт. одной позиции 4BHNN0-0-0-0. По 44-ФЗ в заявку ставится артикул "
        "из извещения. Смесь «1 купол + 49 цилиндров + 1 PTZ» предмету не соответствует. "
        "Эквивалент — только если документация прямо допускает замену.",
        size=10,
        italic=True,
        color=RED,
    )
    add_table(
        doc,
        ["№", "Наименование", "Артикул", "Ед.", "Кол-во", "Цена с НДС, ₽", "Сумма с НДС, ₽"],
        [
            [
                "1",
                "Камера видеонаблюдения цифровая",
                "4BHNN0-0-0-0",
                "шт.",
                "51",
                "22 000",
                "1 122 000",
            ],
        ],
        money_cols={5, 6},
    )
    add_table(
        doc,
        ["Показатель", "Сумма, ₽"],
        [
            ["Всего с НДС по заявке 254397", "1 122 000"],
            ["в том числе НДС 22%", "202 327,87"],
            ["без НДС", "919 672,13"],
        ],
        money_cols={1},
    )
    add_body(
        doc,
        "Артикул в открытых каталогах не найден; 22 000 ₽ — ориентир 4 Мп. Подтвердить у DSSL/НИЦ до подачи. "
        "Если допущен эквивалент: 51× TR-D2151IR3 v2 (R) 2.8 × 16 800 ₽ = 856 800 ₽.",
        size=10,
    )

    add_heading_bar(doc, "Справочно: обе заявки")
    add_table(
        doc,
        ["Комплект", "Сумма с НДС, ₽"],
        [
            ["Заявка 254364 (2 NVR)", "955 480"],
            ["Заявка 254397 (51 камера 4BHNN0)", "1 122 000"],
            ["Сумма двух заявок (не объединять на ЭТП)", "2 077 480"],
        ],
        money_cols={1},
    )

    add_heading_bar(doc, "В. Не входит в извещения")
    add_body(
        doc,
        "Регистратор поставляется без лицензий на камеры и без HDD. Эти позиции нельзя дописать "
        "в ценовой запрос, если их нет в спецификации. Без AnyIP Astra Linux (DSSL 95887) NVR не запишет 51 камеру.",
    )
    add_table(
        doc,
        ["Наименование", "Кол-во", "Цена с НДС, ₽", "Сумма с НДС, ₽"],
        [
            ["Лицензия TRASSIR AnyIP (Astra Linux), DSSL 95887", "51", "5 090", "259 590"],
            ["HDD 10 ТБ SkyHawk AI ST10000VE001", "8", "55 200", "441 600"],
            ["Neuro Detector, доп. каналы (в 2 NVR уже 4)", "47", "6 600", "310 200"],
        ],
        money_cols={2, 3},
    )

    add_heading_bar(doc, "Условия подачи")
    terms = [
        "Подать две ценовые заявки на ЭТП ГПБ до 18.08.2026, 00:00 МСК.",
        "НМЦ не указана. Ориентир участника; демпинг без расчёта себестоимости не закладывать.",
        "Соответствие предмету извещения обязательно (44-ФЗ).",
        "Поставка: 346709, п. Ковалевка, ул. Салютная, д. 2. Срок — по документации процедуры.",
        "Оплата — по 44-ФЗ и контракту ФГКУ, не коммерческая предоплата 70/30.",
        "Цена NVR 477 740 ₽ — открытая розница; у DSSL может быть ниже.",
    ]
    for i, t in enumerate(terms, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(f"{i}. {t}")
        set_run_font(run, size=11)

    add_heading_bar(doc, "Контакты")
    add_table(
        doc,
        ["Поставщик", "Контакт"],
        [
            ["Заказчик", "Пахомов И.В., donsc@dsc.61.mchs.gov.ru, +7 (86350) 2-71-98"],
            ["DSSL / TRASSIR", "8 (800) 100-91-12, dssl.ru, NVR код 80222"],
        ],
    )

    sig = doc.add_paragraph()
    sig.paragraph_format.space_before = Pt(24)
    run = sig.add_run("Поставщик / исполнитель                          Заказчик")
    set_run_font(run, size=11, bold=True)

    sig2 = doc.add_paragraph()
    sig2.paragraph_format.space_before = Pt(18)
    run = sig2.add_run("________________ / ________________              ________________ / ________________")
    set_run_font(run, size=11)
    sig3 = doc.add_paragraph()
    run = sig3.add_run("подпись, ФИО, дата                                         подпись, ФИО, дата")
    set_run_font(run, size=9, color=MUTED)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
