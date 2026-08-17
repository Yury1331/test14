#!/usr/bin/env python3
"""Генерация Excel-спецификации запроса по счетам ЦБ-1548 и ЦБ-1551."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ITEMS = [
    {
        "name": "IP-видеорегистратор",
        "desc": "DHI-NVR5864-EI2, 64-канальный сетевой видеорегистратор WizSense, 2U, 8×SATA до 20 ТБ, "
        "разрешение до 32 Мп, HDMI 8K, AI (AcuPick, SMD Plus, периметр), 2×GbE, ONVIF",
        "vendor": "Dahua",
        "pn": "DHI-NVR5864-EI2",
        "qty": 117,
        "unit": "шт.",
        "price": 49081.31,
    },
    {
        "name": "IP-камера купольная",
        "desc": "DH-IPC-HDBW5442EP-ZE-S3, 4 Мп WizMind, motor-zoom 2.7–12 мм, WDR 140 dB, "
        "ИК до 60 м, IP67/IK10, microSD, PoE, SMD 3.0",
        "vendor": "Dahua",
        "pn": "DH-IPC-HDBW5442EP-ZE-S3",
        "qty": 1242,
        "unit": "шт.",
        "price": 12780.00,
    },
    {
        "name": "IP-камера купольная",
        "desc": "DH-IPC-HDBW5559RP-ASE-IL-0280B, 5 Мп Smart Dual Light WizMind, объектив 2.8 мм, "
        "полноцвет/ИК, IP67, PoE, встроенный микрофон, AI-аналитика",
        "vendor": "Dahua",
        "pn": "DH-IPC-HDBW5559RP-ASE-IL-0280B",
        "qty": 3530,
        "unit": "шт.",
        "price": 11603.00,
    },
]

HEADER_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        cell.border = BORDER


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Спецификация"

    headers = [
        "№",
        "Наименование",
        "Описание",
        "Производитель",
        "PN (парт-номер) / Артикул",
        "ECCN",
        "EAR",
        "Лицензия BIS\n(при закупке в интересах Гос. Заказчика)",
        "Массо-габаритные характеристики:\nвес (кг), ширина (м), длина (м), высота (м)",
        "Кол-во",
        "Ед.изм.",
        "Цена за ед., руб. без НДС",
        "ИТОГО, руб. без НДС",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    total = 0.0
    for i, item in enumerate(ITEMS, 1):
        line_total = item["qty"] * item["price"]
        total += line_total
        ws.append([
            i,
            item["name"],
            item["desc"],
            item["vendor"],
            item["pn"],
            "",
            "",
            "",
            "",
            item["qty"],
            item["unit"],
            item["price"],
            line_total,
        ])

    for row in range(2, 2 + len(ITEMS)):
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = BORDER
            cell.alignment = WRAP
        ws.cell(row=row, column=12).number_format = "#,##0.00"
        ws.cell(row=row, column=13).number_format = "#,##0.00"

    footer_row = 2 + len(ITEMS) + 1
    ws.cell(row=footer_row, column=12, value="Итого, руб. без НДС:").font = Font(bold=True)
    ws.cell(row=footer_row, column=13, value=total).font = Font(bold=True)
    ws.cell(row=footer_row, column=13).number_format = "#,##0.00"

    info_row = footer_row + 2
    info = [
        "Основные требования к товару и условия поставки (заполняется заказчиком):",
        "1) Срок доставки — не более 14 к.д. с момента получения предоплаты;",
        "2) Адрес доставки — ____________ обл., г. ___________, ул. _______________;",
        "3) Способ доставки (авиа или ж/д) — _______________________;",
        "4) Возможность поставки по гарантийному письму до подписания договора: нет;",
        "5) Юридическое наименование поставщика: «………………………………………»;",
        "6) Номер зарегистрированной под МТС сделки у производителя: нет;",
        "7) Условия по оплате: 70 % предоплата, 30 % в течение 3 дней с момента поставки;",
        "8) Аналоги: допускаются (уточнить исключения при необходимости);",
        "9) Спецификацию подготовил: ФИО сотрудника;",
        "10) Тип договора (разовый/рамочный): разовый;",
        "11) Технические контакты от заказчика: ФИО, E-mail, тел.;",
        "12) Контакты от поставщика: ФИО, E-mail, тел.;",
        "",
        "Заказчик: ПАО «МТС», ИНН 7740000076, КПП 770901001",
        "Основание: счета-договоры № ЦБ-1548 и № ЦБ-1551 от 20.07.2026 (ООО «ЭФФОРТ ТЕЛЕКОМ»)",
        "",
        "С текстом Заверения ознакомлен, согласие на подписание подтверждаю. "
        "Товары, поставляемые по настоящему Коммерческому предложению, не подпадают "
        "под действие требований Законодательства в области экспортного контроля.",
    ]
    for j, line in enumerate(info):
        ws.cell(row=info_row + j, column=1, value=line)
        ws.merge_cells(start_row=info_row + j, start_column=1, end_row=info_row + j, end_column=13)
        ws.cell(row=info_row + j, column=1).alignment = WRAP

    widths = [5, 18, 42, 12, 28, 8, 8, 14, 18, 8, 8, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    out = "/workspace/Подбор оборудования/спецификация-запрос-цб-1548-1551.xlsx"
    wb.save(out)
    print(f"Saved: {out}")
    print(f"Total without VAT: {total:,.2f}")


if __name__ == "__main__":
    main()
